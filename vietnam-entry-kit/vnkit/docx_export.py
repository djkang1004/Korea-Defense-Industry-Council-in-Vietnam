"""마크다운 보고서 → Word(.docx) 변환 (제목·표·목록·인용 지원).

최종 고객 제출본은 SMBL 레터지 서식(smbl-market-report 스킬)으로 다시 앉히는 것을 권장한다.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _runs(par, text: str):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            par.add_run(text[pos:m.start()].replace("`", ""))
        par.add_run(m.group(1)).bold = True
        pos = m.end()
    if pos < len(text):
        par.add_run(text[pos:].replace("`", ""))


def _set_font(doc: Document):
    st = doc.styles["Normal"]
    st.font.name = "Malgun Gothic"
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")


def md_to_docx(md: str, dest: Path) -> Path:
    doc = Document()
    _set_font(doc)
    lines = md.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i].strip())
                i += 1
            rows = [[c.strip() for c in r.strip("|").split("|")] for r in block if not re.match(r"^\|[\s\-|:]+\|$", r)]
            if rows:
                ncol = max(len(r) for r in rows)
                t = doc.add_table(rows=len(rows), cols=ncol)
                t.style = "Table Grid"
                t.alignment = WD_TABLE_ALIGNMENT.CENTER
                for ri, r in enumerate(rows):
                    for ci in range(ncol):
                        cell = t.cell(ri, ci)
                        cell.text = ""
                        _runs(cell.paragraphs[0], r[ci] if ci < len(r) else "")
                        for run in cell.paragraphs[0].runs:
                            run.font.size = Pt(8.5)
                            if ri == 0:
                                run.bold = True
                doc.add_paragraph()
            continue
        if line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 4)
            h = doc.add_heading(line.lstrip("#").strip(), level=max(0, level - 1) if level > 1 else 0)
            for run in h.runs:
                run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
        elif line.startswith("> "):
            p = doc.add_paragraph(style="Intense Quote")
            _runs(p, line[2:])
        elif re.match(r"^\s*[-*] ", line):
            _runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*] ", "", line))
        elif re.match(r"^\d+\. ", line):
            _runs(doc.add_paragraph(style="List Number"), re.sub(r"^\d+\. ", "", line))
        elif line.strip():
            _runs(doc.add_paragraph(), line)
        i += 1
    doc.save(dest)
    return dest
